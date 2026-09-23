"""Проверка экспорта: .venv/bin/python scripts/test_export.py [out_dir]
Генерирует DOCX/PDF (обычный и обезличенный) из scripts/export_fixture.json и делает базовые проверки."""
import io
import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.export.docx_export import build_docx  # noqa: E402
from app.export.pdf_export import build_pdf  # noqa: E402

here = os.path.dirname(os.path.abspath(__file__))
meeting = json.load(open(os.path.join(here, "export_fixture.json"), encoding="utf-8"))
out = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="hatshy_export_")
os.makedirs(out, exist_ok=True)
ALL = {"summary", "tasks", "transcript", "people", "sign"}
real_names = [s["name"] for s in meeting["speakers"]]

for anon in (False, True):
    sfx = "_anon" if anon else ""
    pdf = build_pdf(meeting, ALL, anon)
    assert pdf[:5] == b"%PDF-", "not a PDF"
    open(os.path.join(out, f"protocol{sfx}.pdf"), "wb").write(pdf)

    data = build_docx(meeting, ALL, anon)
    open(os.path.join(out, f"protocol{sfx}.docx"), "wb").write(data)
    import docx
    doc = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert len(doc.tables) == 2, f"expected 2 tables, got {len(doc.tables)}"
    assert len(doc.tables[0].rows) == len(meeting["tasks"]) + 1
    assert "Протокол совещания" in text and "Стенограмма" in text and "Председатель" in text
    leaked = [n for n in real_names if n in text or n.split()[0] in text]
    if anon:
        assert not leaked, f"names leaked: {leaked}"
        assert "Участник 1" in text
    else:
        assert len(leaked) == len(real_names)
    print(f"anon={anon}: pdf {len(pdf)} B, docx {len(data)} B, tables={len(doc.tables)}, ok")

# подмножество разделов
d = build_docx(meeting, {"tasks"}, False)
doc = docx.Document(io.BytesIO(d))
assert len(doc.tables) == 1 and not any("Стенограмма" in p.text for p in doc.paragraphs)
p = build_pdf(meeting, {"tasks"}, False)
assert p[:5] == b"%PDF-"
# пустое совещание
empty = {"id": 9, "title": "Пусто", "date": "2026-09-23", "segments": [], "tasks": [], "speakers": [], "summary": None}
build_pdf(empty, None, True)
build_docx(empty, None, True)
print("sections subset + empty meeting ok")
print("out:", out)
