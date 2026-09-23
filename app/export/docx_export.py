"""Экспорт протокола в DOCX (python-docx). Шрифт задаётся по имени (Arial) для всех диапазонов,
тематические шрифты заголовков убраны, чтобы кириллица и казахские буквы шли тем же шрифтом."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .common import prepare

FONT = "Arial"
INK = RGBColor(0x1C, 0x20, 0x28)
MUTED = RGBColor(0x69, 0x70, 0x7D)
ACCENT = RGBColor(0x26, 0x46, 0x8C)
HEAD_FILL = "E8EDF6"


def _set_rfonts(el, font=FONT):
    """Жёстко прописывает шрифт во все диапазоны и убирает theme-атрибуты."""
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), font)


def _setup_styles(doc):
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    _set_rfonts(normal.element)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Heading 1", 13, ACCENT), ("Heading 2", 11.5, INK), ("Title", 20, INK)):
        s = st[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = color
        _set_rfonts(s.element)
        s.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        s.paragraph_format.space_after = Pt(4)
    # rPrDefault документа тоже часто ссылается на тему
    rpr_default = doc.styles.element.find(qn("w:docDefaults"))
    if rpr_default is not None:
        for rf in rpr_default.iter(qn("w:rFonts")):
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)


def _run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def _table(doc, headers, widths_mm, aligns):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # ширины колонок в сетке таблицы (их читают Word/LibreOffice/Quick Look)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_mm):
        gc.set(qn("w:w"), str(Mm(w).twips))
    tblpr = t._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(Mm(sum(widths_mm)).twips))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        if edge in ("left", "right", "insideV"):
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "6")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "C8CDD7")
        borders.append(e)
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    tblpr.append(borders)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Mm(widths_mm[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = aligns[i]
        _run(p, h, bold=True, size=9.5)
        _shade(c, HEAD_FILL)
    _repeat_header(t.rows[0])
    return t


def _add_row(t, widths_mm, aligns):
    row = t.add_row()
    _no_split(row)
    for i, c in enumerate(row.cells):
        c.width = Mm(widths_mm[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        c.paragraphs[0].alignment = aligns[i]
        c.paragraphs[0].paragraph_format.space_after = Pt(1)
    return row.cells


def _page_number_footer(section, text):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _run(p, text + "    ·    стр. ", size=7.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    sub = _run(p, "1", size=7.5, color=MUTED)._r  # run с форматированием переносим внутрь поля
    p._p.remove(sub)
    fld.append(sub)
    p._p.append(fld)


def _bullets(doc, items, numbered=False):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(6)
        p.paragraph_format.first_line_indent = Mm(-5)
        p.paragraph_format.space_after = Pt(2)
        _run(p, f"{i}. " if numbered else "•  ")
        _run(p, it)


def build_docx(meeting: dict, sections, anon: bool) -> bytes:
    d = prepare(meeting, sections, anon)
    sec = d["sections"]
    doc = Document()
    _setup_styles(doc)
    s0 = doc.sections[0]
    s0.page_width, s0.page_height = Mm(210), Mm(297)
    s0.left_margin = s0.right_margin = Mm(20)
    s0.top_margin, s0.bottom_margin = Mm(18), Mm(18)
    _page_number_footer(s0, d["footer"])
    cp = doc.core_properties
    cp.title, cp.author, cp.subject = "Протокол совещания", "AI Hatshy", d["topic"]

    # --- шапка
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    _run(p, d["title"], bold=True, size=20)
    if d["org"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, d["org"], italic=True, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    _run(p, "Тема: ", bold=True, size=12)
    _run(p, d["topic"], bold=True, size=12)
    if d["meta"]:
        p = doc.add_paragraph()
        for i, (k, v) in enumerate(d["meta"]):
            if i:
                _run(p, "   ·   ", size=9.5, color=MUTED)
            _run(p, f"{k}: ", bold=True, size=9.5, color=MUTED)
            _run(p, v, size=9.5, color=MUTED)
    if d["participants"]:
        p = doc.add_paragraph()
        _run(p, f"Участники ({len(d['participants'])}): ", bold=True, size=9.5, color=MUTED)
        _run(p, ", ".join(d["participants"]), size=9.5, color=MUTED)
    if d["anon"]:
        p = doc.add_paragraph()
        _run(p, "Персональные данные обезличены: ФИО заменены на «Участник N».", italic=True, size=8.5, color=MUTED)

    # --- краткое содержание и решения
    if "summary" in sec:
        doc.add_heading("Краткое содержание", level=1)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _run(p, d["summary"] or "Краткое содержание не сформировано.")
        if d["topics"]:
            p = doc.add_paragraph()
            _run(p, "Обсуждавшиеся вопросы:", bold=True)
            _bullets(doc, [f"{lbl} ({t})" if t else lbl for lbl, t in d["topics"]], numbered=True)
        if d["numbers"]:
            p = doc.add_paragraph()
            _run(p, "Ключевые показатели:", bold=True)
            _bullets(doc, [f"{v} — {m}" if m else v for v, m in d["numbers"]])
        doc.add_heading("Принятые решения", level=1)
        if d["decisions"]:
            _bullets(doc, d["decisions"], numbered=True)
        else:
            _run(doc.add_paragraph(), "Отдельные решения не зафиксированы.", color=MUTED)

    # --- поручения
    if "tasks" in sec:
        doc.add_heading("Поручения", level=1)
        if d["tasks"]:
            W = (10, 90, 43, 27)
            AL = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER)
            t = _table(doc, ("№", "Поручение", "Ответственный", "Срок"), W, AL)
            for task in d["tasks"]:
                c = _add_row(t, W, AL)
                _run(c[0].paragraphs[0], task["n"], size=9.5)
                pp = c[1].paragraphs[0]
                _run(pp, task["title"], size=9.5)
                if task["done"]:
                    _run(pp, "  [выполнено]", bold=True, size=8.5, color=ACCENT)
                if task["quote"]:
                    q = c[1].add_paragraph()
                    q.paragraph_format.space_after = Pt(1)
                    _run(q, f"«{task['quote']}»" + (f" [{task['t']}]" if task["t"] else ""),
                         italic=True, size=8, color=MUTED)
                _run(c[2].paragraphs[0], task["owner"], bold=True, size=9.5)
                if task["co_owners"]:
                    _run(c[2].add_paragraph(), "соисполн.: " + ", ".join(task["co_owners"]), size=8.5, color=MUTED)
                _run(c[3].paragraphs[0], task["due"], size=9.5)
        else:
            _run(doc.add_paragraph(), "Поручения не выявлены.", color=MUTED)

    # --- стенограмма
    if "transcript" in sec:
        doc.add_heading("Стенограмма", level=1)
        if d["transcript"]:
            for seg in d["transcript"]:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                _run(p, f"[{seg['t']}] ", bold=True, size=9.5, color=MUTED)
                _run(p, seg["name"], bold=True, size=9.5)
                _run(p, (f" ({seg['lang']})" if seg["lang"] else "") + ": ", size=9.5, color=MUTED)
                _run(p, seg["text"], size=9.5)
        else:
            _run(doc.add_paragraph(), "Стенограмма отсутствует.", color=MUTED)

    # --- участники
    if "people" in sec and d["people"]:
        doc.add_heading("Участники", level=1)
        W = (10, 55, 80, 25)
        AL = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER)
        t = _table(doc, ("№", "Участник", "Роль / должность", "Доля речи"), W, AL)
        for i, person in enumerate(d["people"], 1):
            c = _add_row(t, W, AL)
            for j, v in enumerate((str(i), person["name"], person["role"], person["pct"])):
                _run(c[j].paragraphs[0], v, size=9.5)

    # --- подписи
    if "sign" in sec:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(28)
        p.paragraph_format.keep_with_next = True
        _run(p, "Председатель  ______________________" + (f"  / {d['chair']} /" if d["chair"] else ""))
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        _run(p, "Секретарь  ______________________")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
